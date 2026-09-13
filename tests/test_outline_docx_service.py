import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.services.outline_docx_service import create_outline_docx


class OutlineDocxServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_creates_docx_with_standard_data(self):
        sample_data = {
            "tutorial_name": "Introduction to Docker",
            "outline_type": "FOSS",
            "target_audience": "Software Developers",
            "entry_behaviour": "Basic command line knowledge",
            "purpose": "Learn containerization basics",
            "recommended_no_of_tutorials": 3,
            "prepared_by": "Course Instructor",
            "date": "2025-01-15",
            "keywords": ["Docker", "Containers", "DevOps"],
            "course_objectives": ["Build Docker images", "Run containers"],
            "topics_included": ["Images", "Containers", "Docker Compose"],
            "topics_not_included": ["Kubernetes"],
            "core_example": "Web application container",
            "allied_examples": ["Database container", "Worker container"],
            "tutorial_rows": [
                {
                    "tutorial_number": 1,
                    "title": "Installing Docker",
                    "topics_details": ["Prerequisites", "Installation steps", "Verification"],
                    "time_seconds": 240,
                    "comments": "Beginner friendly",
                }
            ],
        }

        output_file = Path(self.temp_dir.name) / "test_docker_outline.docx"
        result_path = create_outline_docx(sample_data, str(output_file))

        self.assertTrue(os.path.exists(result_path))
        doc = Document(result_path)

        headings = [p.text for p in doc.paragraphs if p.text.startswith("Course Outline")]
        self.assertTrue(len(headings) > 0)

        # Verify metadata table exists
        self.assertTrue(len(doc.tables) >= 2)
        metadata_table = doc.tables[0]
        self.assertIn("FOSS Version", [row.cells[0].text for row in metadata_table.rows])

    def test_creates_docx_with_ict_outline_type(self):
        sample_data = {
            "outline_name": "ICT in Education",
            "outline_type": "ICT",
            "platform_name": "Google Classroom",
            "course_objectives": ["Using digital tools for teaching"],
            "tutorial_rows": [],
        }

        output_file = Path(self.temp_dir.name) / "test_ict_outline.docx"
        result_path = create_outline_docx(sample_data, str(output_file))

        self.assertTrue(os.path.exists(result_path))
        doc = Document(result_path)
        metadata_table = doc.tables[0]
        labels = [row.cells[0].text for row in metadata_table.rows]
        self.assertIn("ICT Platform/Program", labels)

    def test_handles_minimal_empty_outline_gracefully(self):
        minimal_data = {}
        output_file = Path(self.temp_dir.name) / "test_empty_outline.docx"
        result_path = create_outline_docx(minimal_data, str(output_file))

        self.assertTrue(os.path.exists(result_path))
        doc = Document(result_path)
        self.assertTrue(len(doc.tables) >= 1)


if __name__ == "__main__":
    unittest.main()
