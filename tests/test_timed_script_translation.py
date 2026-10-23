"""Unit tests for timed script parsing, multi-row Sarvam translation, and DOCX export."""
import asyncio
import io
import json
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch
from docx import Document
import httpx

from src.services.timed_script_parser import (
    parse_docx_timed_script,
    parse_json_timed_script,
    parse_txt_timed_script,
    parse_timed_script_file,
)
from src.services.timed_script_translation_service import (
    _pack_rows_into_batches,
    _translate_batch_sarvam,
    translate_timed_script,
    batch_translate_timed_script,
    export_translated_timed_script_docx,
    TimedScriptTranslationResult,
)


class TimedScriptTranslationUnitTests(unittest.IsolatedAsyncioTestCase):

    def test_parse_docx_timed_script(self):
        """Test parsing of a 2-column DOCX table with Time and Narration."""
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Time"
        table.rows[0].cells[1].text = "Narration"

        r1 = table.add_row().cells
        r1[0].text = "00:00"
        r1[1].text = "Welcome to the tutorial."

        r2 = table.add_row().cells
        r2[0].text = "00:05"
        r2[1].text = "Next step is here."

        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        rows = parse_docx_timed_script(file_bytes)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["time"], "00:00")
        self.assertEqual(rows[0]["narration"], "Welcome to the tutorial.")
        self.assertEqual(rows[1]["time"], "00:05")
        self.assertEqual(rows[1]["narration"], "Next step is here.")

    def test_parse_json_timed_script_variants(self):
        """Test parsing JSON with either 'rows' or 'sentences' (from Whisper)."""
        whisper_json = {
            "sentences": [
                {"time_range": "00:00 - 00:04", "text": "First sentence."},
                {"time_range": "00:04 - 00:08", "text": "Second sentence."}
            ]
        }
        rows = parse_json_timed_script(json.dumps(whisper_json).encode())
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["time"], "00:00 - 00:04")
        self.assertEqual(rows[0]["narration"], "First sentence.")

        custom_json = {
            "rows": [
                {"time": "01:00", "narration": "Row one."}
            ]
        }
        rows2 = parse_json_timed_script(json.dumps(custom_json).encode())
        self.assertEqual(len(rows2), 1)
        self.assertEqual(rows2[0]["time"], "01:00")
        self.assertEqual(rows2[0]["narration"], "Row one.")

    def test_parse_txt_timed_script(self):
        """Test parsing plain text lines with timestamps."""
        txt = "00:00 Welcome everyone\n00:07 Let's begin"
        rows = parse_txt_timed_script(txt.encode())
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["time"], "00:00")
        self.assertEqual(rows[0]["narration"], "Welcome everyone")

    def test_pack_rows_into_batches(self):
        """Packing rows should respect max_chars limits."""
        rows = [
            {"row_number": 1, "time": "00:00", "narration": "A" * 100},
            {"row_number": 2, "time": "00:05", "narration": "B" * 100},
            {"row_number": 3, "time": "00:10", "narration": "C" * 100},
        ]
        # max_chars=150 should split rows into 3 separate batches
        batches = _pack_rows_into_batches(rows, max_chars=150)
        self.assertEqual(len(batches), 3)

        # max_chars=500 should fit all into 1 batch
        batches_single = _pack_rows_into_batches(rows, max_chars=500)
        self.assertEqual(len(batches_single), 1)
        self.assertEqual(len(batches_single[0]), 3)

    @patch("src.services.timed_script_translation_service.httpx.AsyncClient.post")
    async def test_translate_batch_sarvam_success(self, mock_post):
        """Test multi-row batch translation with Sarvam."""
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.json.return_value = {
            "translated_text": "[#1] अनुवाद १\n[#2] अनुवाद २"
        }
        mock_post.return_value = mock_response

        batch = [
            {"row_number": 1, "time": "00:00", "narration": "English 1"},
            {"row_number": 2, "time": "00:05", "narration": "English 2"},
        ]

        async with httpx.AsyncClient() as client:
            res = await _translate_batch_sarvam(
                client=client,
                batch=batch,
                target_bcp47="hi-IN",
                api_key="test-key",
            )
            self.assertEqual(res[1], "अनुवाद १")
            self.assertEqual(res[2], "अनुवाद २")

    @patch("src.services.timed_script_translation_service._translate_batch_sarvam", new_callable=AsyncMock)
    async def test_translate_timed_script_sarvam_full(self, mock_batch_translate):
        """translate_timed_script translates rows and maps translations back."""
        mock_batch_translate.return_value = {1: "नमस्ते", 2: "आगे बढ़ें"}

        rows = [
            {"row_number": 1, "time": "00:00", "narration": "Hello"},
            {"row_number": 2, "time": "00:05", "narration": "Continue"}
        ]

        with patch.dict("os.environ", {"SARVAM_API_KEY": "valid-key", "TRANSLATION_PROVIDER": "sarvam"}):
            res = await translate_timed_script(rows, "hi", title="Test Tutorial")

        self.assertTrue(res.success)
        self.assertEqual(res.language_code, "hi")
        self.assertEqual(len(res.rows), 2)
        self.assertEqual(res.rows[0]["narration_translated"], "नमस्ते")
        self.assertEqual(res.rows[0]["narration_hi"], "नमस्ते")
        self.assertEqual(res.rows[1]["narration_translated"], "आगे बढ़ें")

    @patch("src.services.timed_script_translation_service._translate_timed_script_gemini", new_callable=AsyncMock)
    @patch("src.services.timed_script_translation_service._translate_batch_sarvam", new_callable=AsyncMock)
    async def test_translate_timed_script_fallback_to_gemini(self, mock_sarvam, mock_gemini):
        """When Sarvam fails, translate_timed_script falls back to Gemini."""
        mock_sarvam.side_effect = RuntimeError("Sarvam 500 Internal Server Error")
        mock_gemini.return_value = {1: "मिथुन अनुवाद"}

        rows = [{"row_number": 1, "time": "00:00", "narration": "Hello"}]

        with patch.dict("os.environ", {
            "SARVAM_API_KEY": "valid-key",
            "GEMINI_API_KEY": "valid-gemini-key",
            "TRANSLATION_PROVIDER": "sarvam"
        }):
            res = await translate_timed_script(rows, "hi")

        self.assertTrue(res.success)
        self.assertEqual(res.rows[0]["narration_translated"], "मिथुन अनुवाद")
        mock_sarvam.assert_awaited_once()
        mock_gemini.assert_awaited_once()

    @patch("src.services.timed_script_translation_service.translate_timed_script", new_callable=AsyncMock)
    async def test_batch_translate_timed_script(self, mock_translate):
        """batch_translate_timed_script translates across multiple requested languages."""
        async def fake_translate(rows, lang, **kwargs):
            return TimedScriptTranslationResult(
                language="Lang",
                language_code=lang,
                language_native="Native",
                rows=rows,
                total_rows=len(rows),
                success=True
            )
        mock_translate.side_effect = fake_translate

        rows = [{"row_number": 1, "time": "00:00", "narration": "Hello"}]
        results = await batch_translate_timed_script(rows, ["hi", "ta"])

        self.assertEqual(len(results), 2)
        self.assertTrue(all(r.success for r in results))
        self.assertEqual([r.language_code for r in results], ["hi", "ta"])

    def test_export_translated_timed_script_docx(self):
        """Test DOCX export generating proper Spoken Tutorial table layout."""
        rows = [
            {"row_number": 1, "time": "00:00", "narration": "Hello", "narration_translated": "नमस्ते"},
            {"row_number": 2, "time": "00:07", "narration": "World", "narration_translated": "दुनिया"}
        ]
        buffer = export_translated_timed_script_docx(
            rows=rows,
            language_code="hi",
            language_name="Hindi",
            language_native="हिंदी",
            title="Spoken Tutorial",
            total_duration="00:07"
        )
        self.assertIsInstance(buffer, io.BytesIO)
        doc = Document(buffer)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)  # Header + 2 rows
        self.assertEqual(table.rows[0].cells[0].text, "Time")
        self.assertEqual(table.rows[0].cells[1].text, "Narration")
        self.assertEqual(table.rows[1].cells[0].text, "00:00")
        self.assertEqual(table.rows[1].cells[1].text, "नमस्ते")
        self.assertEqual(table.rows[2].cells[0].text, "00:07")
        self.assertEqual(table.rows[2].cells[1].text, "दुनिया")

    @patch("src.api.routes.timed_script.batch_translate_timed_script", new_callable=AsyncMock)
    async def test_batch_translate_endpoint(self, mock_batch_translate):
        """Test POST /timed-script/batch-translate endpoint."""
        from src.api.routes.timed_script import (
            batch_translate_timed_script_endpoint,
            BatchTranslateTimedScriptRequest,
        )

        mock_batch_translate.return_value = [
            TimedScriptTranslationResult(
                language="Hindi",
                language_code="hi",
                language_native="हिंदी",
                title="Test",
                total_duration="00:05",
                total_rows=1,
                rows=[{"row_number": 1, "time": "00:00", "narration": "Hello", "narration_translated": "नमस्ते"}],
                success=True
            )
        ]

        req = BatchTranslateTimedScriptRequest(
            rows=[{"row_number": 1, "time": "00:00", "narration": "Hello"}],
            languages=["hi"],
            title="Test",
            total_duration="00:05"
        )

        user = SimpleNamespace(sub="user-123", email="user@edupyramids.org")
        res = await batch_translate_timed_script_endpoint(req, current_user=user)

        self.assertEqual(res.total_success, 1)
        self.assertEqual(res.results[0].language_code, "hi")

    async def test_export_translated_docx_endpoint(self):
        """Test POST /timed-script/export-translated-docx endpoint."""
        from src.api.routes.timed_script import (
            export_translated_docx_endpoint,
            ExportTranslatedTimedScriptRequest,
        )

        req = ExportTranslatedTimedScriptRequest(
            rows=[{"row_number": 1, "time": "00:00", "narration": "Hello", "narration_translated": "नमस्ते"}],
            language_code="hi",
            language_name="Hindi",
            language_native="हिंदी",
            title="Calcium Tutorial",
            total_duration="00:05"
        )
        user = SimpleNamespace(sub="user-123", email="user@edupyramids.org")
        response = await export_translated_docx_endpoint(req, current_user=user)

        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIn("Calcium_Tutorial_hi_timed_script.docx", response.headers["Content-Disposition"])

    @patch("src.api.routes.timed_script.batch_translate_timed_script", new_callable=AsyncMock)
    @patch("src.api.routes.timed_script.parse_timed_script_file")
    async def test_translate_file_endpoint_language_parsing(self, mock_parse, mock_batch_translate):
        """translate_timed_script_file_endpoint correctly parses JSON array, names, and strings."""
        from src.api.routes.timed_script import translate_timed_script_file_endpoint

        mock_parse.return_value = {
            "title": "Tutorial",
            "total_duration": "01:00",
            "rows": [{"row_number": 1, "time": "00:00", "narration": "Hello"}]
        }
        mock_batch_translate.return_value = [
            TimedScriptTranslationResult(
                language="Hindi",
                language_code="hi",
                language_native="हिंदी",
                title="Tutorial",
                total_duration="01:00",
                total_rows=1,
                rows=[{"row_number": 1, "time": "00:00", "narration": "Hello", "narration_translated": "नमस्ते"}],
                success=True
            )
        ]

        class DummyUploadFile:
            def __init__(self, filename, content):
                self.filename = filename
                self.content = content
            async def read(self):
                return self.content

        user = SimpleNamespace(sub="user-123", email="user@edupyramids.org")

        # Test case 1: JSON array string '["hi"]'
        f1 = DummyUploadFile("timed.docx", b"dummy")
        res1 = await translate_timed_script_file_endpoint(file=f1, languages='["hi"]', current_user=user)
        self.assertEqual(res1.total_success, 1)

        # Test case 2: Language name "Hindi"
        f2 = DummyUploadFile("timed.docx", b"dummy")
        res2 = await translate_timed_script_file_endpoint(file=f2, languages='Hindi', current_user=user)
        self.assertEqual(res2.total_success, 1)

        # Test case 3: Double JSON encoded string '"[\"hi\"]"'
        f3 = DummyUploadFile("timed.docx", b"dummy")
        res3 = await translate_timed_script_file_endpoint(file=f3, languages='"[\"hi\"]"', current_user=user)
        self.assertEqual(res3.total_success, 1)


if __name__ == "__main__":
    unittest.main()
