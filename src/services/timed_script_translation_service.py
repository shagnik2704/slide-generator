"""
Timed Script Translation Service.
Translates time-aligned narration scripts to Indian languages using Sarvam AI (sarvam-translate:v1)
with multi-row batching and Google Gemini fallback.
"""
import asyncio
import io
import json
import logging
import os
import re
import time
from typing import Any, Dict, List, Optional
import httpx
from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from src.services.translation_service import (
    DEFAULT_SARVAM_TRANSLATE_MODEL,
    SARVAM_MAX_RETRIES,
    SARVAM_TRANSLATE_URL,
    resolve_language_info,
)

logger = logging.getLogger(__name__)

# Max batch character budget (safely under Sarvam's 2,000 char limit)
BATCH_CHAR_LIMIT = 1200
BATCH_CONCURRENCY = 10
PIPE_TAG_PATTERN = re.compile(r'(?:^|\n)\s*(\d+)\s*\|\|\|\s*(.*?)(?=\n\s*\d+\s*\|\|\||$)', re.DOTALL)
ROW_TAG_PATTERN = re.compile(r'\[#?(\d+)\]\s*(.*?)(?=\[#?\d+\]|$)', re.DOTALL)


class TimedScriptTranslationResult(BaseModel):
    """Result of translating a timed script to a single language."""
    language: str
    language_code: str
    language_native: str
    title: str = "Timed Script"
    total_duration: str = "00:00"
    total_rows: int = 0
    rows: List[Dict[str, Any]] = Field(default_factory=list)
    success: bool = True
    error: Optional[str] = None


class BatchTimedScriptTranslationResponse(BaseModel):
    """Response containing multiple translated timed scripts."""
    results: List[TimedScriptTranslationResult]
    total_requested: int
    total_success: int


def _pack_rows_into_batches(rows: List[Dict[str, Any]], max_chars: int = BATCH_CHAR_LIMIT) -> List[List[Dict[str, Any]]]:
    """
    Pack rows into multi-row batches such that the combined text of each batch
    remains within max_chars.
    """
    batches: List[List[Dict[str, Any]]] = []
    current_batch: List[Dict[str, Any]] = []
    current_len = 0

    for row in rows:
        row_num = row.get("row_number", len(current_batch) + 1)
        narration = row.get("narration", "")
        line_len = len(f"{row_num} ||| {narration}\n")

        if current_batch and (current_len + line_len > max_chars):
            batches.append(current_batch)
            current_batch = [row]
            current_len = line_len
        else:
            current_batch.append(row)
            current_len += line_len

    if current_batch:
        batches.append(current_batch)

    return batches


async def _translate_batch_sarvam(
    client: httpx.AsyncClient,
    batch: List[Dict[str, Any]],
    target_bcp47: str,
    api_key: str,
    model: str = DEFAULT_SARVAM_TRANSLATE_MODEL,
    source_code: str = "en-IN",
) -> Dict[int, str]:
    """
    Translate a single multi-row batch formatted with 'N ||| narration' tags.
    Returns a dict mapping row_number -> translated_text.
    """
    # Build tagged input using pipe delimiter (leaves tags intact in Sarvam)
    lines = [f"{r.get('row_number')} ||| {r.get('narration', '')}" for r in batch]
    input_text = "\n".join(lines)

    payload = {
        "input": input_text,
        "source_language_code": source_code,
        "target_language_code": target_bcp47,
        "model": model,
    }
    headers = {
        "Content-Type": "application/json",
        "api-subscription-key": api_key,
    }

    last_error = None
    for attempt in range(SARVAM_MAX_RETRIES + 1):
        try:
            response = await client.post(
                SARVAM_TRANSLATE_URL,
                json=payload,
                headers=headers,
            )
            if response.status_code == 200:
                translated_text = response.json().get("translated_text", "")
                results: Dict[int, str] = {}

                # 1. Primary: match 'N ||| translated_text'
                matches = PIPE_TAG_PATTERN.findall(translated_text)
                for num_str, content in matches:
                    try:
                        results[int(num_str)] = " ".join(content.strip().split())
                    except ValueError:
                        continue

                # 2. Secondary fallback: check legacy '[#N]' or '[N]'
                if not results:
                    legacy_matches = ROW_TAG_PATTERN.findall(translated_text)
                    for num_str, content in legacy_matches:
                        try:
                            results[int(num_str)] = " ".join(content.strip().split())
                        except ValueError:
                            continue

                # 3. For any row that was dropped by delimiter parsing, translate individually
                for r in batch:
                    rn = r.get("row_number")
                    if rn not in results or not results[rn]:
                        narration = r.get("narration", "")
                        if not narration:
                            results[rn] = ""
                            continue
                        try:
                            single_resp = await client.post(
                                SARVAM_TRANSLATE_URL,
                                json={
                                    "input": narration,
                                    "source_language_code": source_code,
                                    "target_language_code": target_bcp47,
                                    "model": model,
                                },
                                headers=headers,
                            )
                            if single_resp.status_code == 200:
                                results[rn] = single_resp.json().get("translated_text", "").strip()
                            else:
                                results[rn] = narration
                        except Exception:
                            results[rn] = narration

                return results

            status = response.status_code
            if status != 429 and status < 500:
                err_msg = ""
                try:
                    err_data = response.json()
                    err_msg = err_data.get("message") or err_data.get("detail") or str(err_data)
                except Exception:
                    err_msg = response.text
                raise RuntimeError(f"Sarvam translation error ({status}): {err_msg}")

            if attempt < SARVAM_MAX_RETRIES:
                await asyncio.sleep((attempt + 1) * 1.0)
                continue

            response.raise_for_status()

        except (httpx.RequestError, httpx.TimeoutException) as e:
            last_error = e
            if attempt < SARVAM_MAX_RETRIES:
                await asyncio.sleep((attempt + 1) * 1.0)
                continue
            raise

    if last_error:
        raise last_error
    raise RuntimeError("Sarvam batch translation failed after retries")


async def _translate_timed_script_gemini(
    rows: List[Dict[str, Any]],
    target_code: str,
    lang_name: str,
    lang_native: str,
) -> Dict[int, str]:
    """Fallback translation for timed script rows using Google Gemini."""
    from langchain_google_genai import ChatGoogleGenerativeAI
    from pydantic import BaseModel as PydanticBaseModel

    class TimedRowTranslation(PydanticBaseModel):
        row_number: int
        translated_text: str

    class TimedBatchTranslation(PydanticBaseModel):
        translations: List[TimedRowTranslation]

    llm = ChatGoogleGenerativeAI(
        model="gemini-2.5-flash",
        temperature=0.3,
        api_key=os.getenv("GEMINI_API_KEY"),
    )
    structured_llm = llm.with_structured_output(TimedBatchTranslation)

    prompt = f"""You are an expert English-to-{lang_name} ({lang_native}) translator for educational Spoken Tutorial timed scripts.
Translate the narration of each numbered row into natural {lang_name}.
Keep technical terms (like software commands, filenames, units) intact or transliterated.

Rows to translate:
{json.dumps([{"row_number": r["row_number"], "narration": r["narration"]} for r in rows], indent=2)}

Return the translated text for each row number."""

    result = await structured_llm.ainvoke(prompt)
    if not result or not result.translations:
        raise RuntimeError("Gemini timed script translation returned empty result")

    return {item.row_number: item.translated_text for item in result.translations}


async def translate_timed_script(
    rows: List[Dict[str, Any]],
    target_language: str,
    title: str = "Timed Script",
    total_duration: str = "00:00",
) -> TimedScriptTranslationResult:
    """
    Translate a timed script to a single target language using Sarvam multi-row batching.
    """
    resolved = resolve_language_info(target_language)
    if not resolved:
        return TimedScriptTranslationResult(
            language=target_language,
            language_code=target_language,
            language_native="Unknown",
            title=title,
            total_duration=total_duration,
            total_rows=len(rows),
            rows=[],
            success=False,
            error=f"Unsupported language: {target_language}",
        )

    lang_code, lang_name, lang_native, target_bcp47 = resolved

    if not rows:
        return TimedScriptTranslationResult(
            language=lang_name,
            language_code=lang_code,
            language_native=lang_native,
            title=title,
            total_duration=total_duration,
            total_rows=0,
            rows=[],
            success=False,
            error="No rows found in timed script",
        )

    print(f"⏱️ Translating timed script ({len(rows)} rows) to {lang_name} ({lang_native})...")
    t0 = time.perf_counter()

    provider = os.getenv("TRANSLATION_PROVIDER", "sarvam").lower().strip()
    sarvam_key = (os.getenv("SARVAM_API_KEY") or "").strip()
    sarvam_model = os.getenv("SARVAM_TRANSLATE_MODEL", DEFAULT_SARVAM_TRANSLATE_MODEL).strip()

    translated_map: Dict[int, str] = {}

    try:
        if provider == "sarvam" and sarvam_key:
            try:
                batches = _pack_rows_into_batches(rows, max_chars=BATCH_CHAR_LIMIT)
                semaphore = asyncio.Semaphore(BATCH_CONCURRENCY)
                limits = httpx.Limits(max_keepalive_connections=20, max_connections=50)

                async with httpx.AsyncClient(timeout=30.0, limits=limits) as client:
                    async def process_batch(b):
                        async with semaphore:
                            return await _translate_batch_sarvam(
                                client=client,
                                batch=b,
                                target_bcp47=target_bcp47,
                                api_key=sarvam_key,
                                model=sarvam_model,
                            )

                    batch_tasks = [process_batch(b) for b in batches]
                    batch_results = await asyncio.gather(*batch_tasks)

                for b_res in batch_results:
                    translated_map.update(b_res)

            except Exception as sarvam_err:
                logger.warning(
                    "Sarvam timed script translation failed: %s. Attempting Gemini fallback...",
                    sarvam_err,
                )
                if os.getenv("GEMINI_API_KEY"):
                    translated_map = await _translate_timed_script_gemini(
                        rows=rows,
                        target_code=lang_code,
                        lang_name=lang_name,
                        lang_native=lang_native,
                    )
                else:
                    raise sarvam_err
        else:
            if not sarvam_key:
                logger.info("SARVAM_API_KEY not configured; using Gemini translation.")
            translated_map = await _translate_timed_script_gemini(
                rows=rows,
                target_code=lang_code,
                lang_name=lang_name,
                lang_native=lang_native,
            )

        # Assemble final translated rows
        translated_rows: List[Dict[str, Any]] = []
        for r in rows:
            rn = r.get("row_number")
            trans_text = translated_map.get(rn, r.get("narration", ""))
            row_copy = r.copy()
            row_copy["narration_translated"] = trans_text
            row_copy[f"narration_{lang_code}"] = trans_text
            translated_rows.append(row_copy)

        elapsed = time.perf_counter() - t0
        print(f"   ✓ Translated {len(translated_rows)} timed rows to {lang_name} in {elapsed:.2f}s")

        return TimedScriptTranslationResult(
            language=lang_name,
            language_code=lang_code,
            language_native=lang_native,
            title=title,
            total_duration=total_duration,
            total_rows=len(translated_rows),
            rows=translated_rows,
            success=True,
        )

    except Exception as e:
        logger.exception("Timed script translation failed: %s", e)
        return TimedScriptTranslationResult(
            language=lang_name,
            language_code=lang_code,
            language_native=lang_native,
            title=title,
            total_duration=total_duration,
            total_rows=len(rows),
            rows=[],
            success=False,
            error=str(e),
        )


async def batch_translate_timed_script(
    rows: List[Dict[str, Any]],
    languages: List[str],
    title: str = "Timed Script",
    total_duration: str = "00:00",
) -> List[TimedScriptTranslationResult]:
    """
    Translate a timed script to multiple languages in parallel.
    """
    # Resolve languages to canonical codes (e.g. 'hi', 'ta', 'te')
    valid_languages: List[str] = []
    seen = set()
    for lang in languages:
        info = resolve_language_info(lang)
        if info:
            code = info[0]
            if code not in seen:
                seen.add(code)
                valid_languages.append(code)

    if not valid_languages:
        logger.warning("No valid languages provided to batch_translate_timed_script. Received: %s", languages)
        return [TimedScriptTranslationResult(
            language="Unknown",
            language_code="",
            language_native="",
            title=title,
            total_duration=total_duration,
            total_rows=len(rows),
            rows=[],
            success=False,
            error="No valid languages provided",
        )]

    print(f"🌍 Batch translating timed script to {len(valid_languages)} languages: {valid_languages}")

    tasks = [
        translate_timed_script(rows, lang, title=title, total_duration=total_duration)
        for lang in valid_languages
    ]

    results = await asyncio.gather(*tasks, return_exceptions=True)

    processed_results: List[TimedScriptTranslationResult] = []
    for i, res in enumerate(results):
        if isinstance(res, Exception):
            lang = valid_languages[i]
            info = resolve_language_info(lang)
            name = info[1] if info else "Unknown"
            native = info[2] if info else ""
            processed_results.append(TimedScriptTranslationResult(
                language=name,
                language_code=lang,
                language_native=native,
                title=title,
                total_duration=total_duration,
                total_rows=len(rows),
                rows=[],
                success=False,
                error=str(res),
            ))
        else:
            processed_results.append(res)

    return processed_results


def export_translated_timed_script_docx(
    rows: List[Dict[str, Any]],
    language_code: str,
    language_name: str,
    language_native: str,
    title: str = "Timed Script",
    total_duration: str = "00:00",
) -> io.BytesIO:
    """
    Export translated timed script to a Word document matching the Spoken Tutorial table layout:
    Column 1: Time
    Column 2: Narration (in regional language)
    """
    doc = Document()

    # Title
    heading = doc.add_heading(f"{title} ({language_name})", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if heading.runs:
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    # Info summary
    p = doc.add_paragraph()
    p.add_run(f"Language: ").bold = True
    p.add_run(f"{language_name} ({language_native})\n")
    p.add_run(f"Total Duration: ").bold = True
    p.add_run(f"{total_duration}\n")
    p.add_run(f"Total Sentences: ").bold = True
    p.add_run(f"{len(rows)}")

    doc.add_paragraph()

    # 2-column table
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Time"
    hdr_cells[1].text = "Narration"

    for cell in hdr_cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10.5)

    # Populate rows
    for r in rows:
        row_cells = table.add_row().cells
        time_text = r.get("time", "")
        narration_text = r.get("narration_translated") or r.get(f"narration_{language_code}") or r.get("narration", "")

        row_cells[0].text = time_text
        row_cells[1].text = narration_text

        # Styling
        if row_cells[0].paragraphs and row_cells[0].paragraphs[0].runs:
            row_cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        if row_cells[1].paragraphs and row_cells[1].paragraphs[0].runs:
            row_cells[1].paragraphs[0].runs[0].font.size = Pt(10.5)

    # Set column widths & cell padding
    for row in table.rows:
        row.cells[0].width = Inches(1.5)
        row.cells[1].width = Inches(5.0)

        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                '<w:top w:w="160" w:type="dxa"/>'
                '<w:left w:w="200" w:type="dxa"/>'
                '<w:bottom w:w="160" w:type="dxa"/>'
                '<w:right w:w="200" w:type="dxa"/>'
                '</w:tcMar>'
            )
            tcPr.append(tcMar)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
