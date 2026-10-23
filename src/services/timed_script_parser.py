"""
Timed Script Parser - Parse time-aligned scripts from DOCX, PDF, JSON, and TXT files.
Extracts rows with timestamp and narration text.
"""
import io
import json
import logging
import re
from typing import Any, Dict, List, Optional
from docx import Document

logger = logging.getLogger(__name__)

# Matches timestamps like "00:00", "01:23", "00:00 - 00:07", "01:05:20"
TIMESTAMP_PATTERN = re.compile(
    r'^\s*(\d{1,2}:\d{2}(?::\d{2})?(?:\s*[-–—]\s*\d{1,2}:\d{2}(?::\d{2})?)?)\s*(.*)$',
    re.DOTALL
)

TIME_ONLY_PATTERN = re.compile(r'^\s*\d{1,2}:\d{2}(?::\d{2})?(?:\s*[-–—]\s*\d{1,2}:\d{2}(?::\d{2})?)?\s*$')


def _clean_text(text: Optional[str]) -> str:
    """Clean whitespace, replace multiple newlines with single space."""
    if not text:
        return ""
    # Replace non-breaking spaces and collapse whitespace
    text = text.replace('\xa0', ' ')
    return " ".join(text.split()).strip()


def parse_docx_timed_script(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Parse a DOCX document into timed script rows.
    Checks tables first (standard Spoken Tutorial format), then falls back to paragraphs.
    """
    doc = Document(io.BytesIO(file_bytes))
    rows: List[Dict[str, Any]] = []

    # Check tables
    if doc.tables:
        for table in doc.tables:
            for row_idx, row in enumerate(table.rows):
                if len(row.cells) < 2:
                    continue
                col0 = _clean_text(row.cells[0].text)
                col1 = _clean_text(row.cells[1].text)

                # Skip header row (e.g. "Time", "Narration", "Text")
                if col0.lower() in {"time", "time range", "timestamp", "start time"} and \
                   any(h in col1.lower() for h in {"narration", "text", "script", "sentence"}):
                    continue

                if col0 and col1 and (TIME_ONLY_PATTERN.match(col0) or TIMESTAMP_PATTERN.match(col0)):
                    rows.append({
                        "row_number": len(rows) + 1,
                        "time": col0,
                        "narration": col1
                    })

    # If no table rows found, fallback to line-by-line paragraph parsing
    if not rows:
        for p in doc.paragraphs:
            text = _clean_text(p.text)
            if not text:
                continue
            match = TIMESTAMP_PATTERN.match(text)
            if match:
                time_str = match.group(1).strip()
                narration = _clean_text(match.group(2))
                if narration:
                    rows.append({
                        "row_number": len(rows) + 1,
                        "time": time_str,
                        "narration": narration
                    })

    return rows


def parse_pdf_timed_script(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Parse a PDF document into timed script rows using pdfplumber.
    Supports multi-page tables and plain-text fallback.
    """
    import pdfplumber

    rows: List[Dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if not row or len(row) < 2:
                        continue
                    col0 = _clean_text(row[0])
                    col1 = _clean_text(row[1])

                    # Skip headers
                    if col0.lower() in {"time", "time range", "timestamp"} and \
                       any(h in col1.lower() for h in {"narration", "text", "script"}):
                        continue

                    if col0 and col1 and (TIME_ONLY_PATTERN.match(col0) or TIMESTAMP_PATTERN.match(col0)):
                        rows.append({
                            "row_number": len(rows) + 1,
                            "time": col0,
                            "narration": col1
                        })

            # If no table on page, try text lines
            if not tables or not rows:
                text = page.extract_text() or ""
                for line in text.split("\n"):
                    clean_l = _clean_text(line)
                    if not clean_l:
                        continue
                    m = TIMESTAMP_PATTERN.match(clean_l)
                    if m:
                        time_str = m.group(1).strip()
                        narration = _clean_text(m.group(2))
                        if narration:
                            rows.append({
                                "row_number": len(rows) + 1,
                                "time": time_str,
                                "narration": narration
                            })

    return rows


def parse_json_timed_script(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Parse a JSON document containing either:
    1. 'rows': [{'time': ..., 'narration': ...}]
    2. 'sentences': [{'time_range': ..., 'text': ...}] (from Whisper)
    """
    data = json.loads(file_bytes.decode('utf-8'))
    rows: List[Dict[str, Any]] = []

    if isinstance(data, list):
        items = data
    elif isinstance(data, dict):
        items = data.get("rows") or data.get("sentences") or []
    else:
        items = []

    for item in items:
        time_str = item.get("time") or item.get("time_range") or item.get("start_time") or "00:00"
        narration = item.get("narration") or item.get("text") or ""
        if narration:
            rows.append({
                "row_number": len(rows) + 1,
                "time": str(time_str).strip(),
                "narration": _clean_text(narration)
            })

    return rows


def parse_txt_timed_script(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Parse plain text lines with timestamps."""
    text = file_bytes.decode('utf-8', errors='replace')
    rows: List[Dict[str, Any]] = []

    for line in text.splitlines():
        clean_l = _clean_text(line)
        if not clean_l:
            continue
        m = TIMESTAMP_PATTERN.match(clean_l)
        if m:
            time_str = m.group(1).strip()
            narration = _clean_text(m.group(2))
            if narration:
                rows.append({
                    "row_number": len(rows) + 1,
                    "time": time_str,
                    "narration": narration
                })

    return rows


def parse_timed_script_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Main parser dispatching by file extension.
    Returns:
    {
        "title": filename_stem,
        "total_rows": int,
        "total_duration": str,
        "rows": [{"row_number": int, "time": str, "narration": str}]
    }
    """
    fn = filename.lower()
    rows: List[Dict[str, Any]] = []

    if fn.endswith('.docx') or fn.endswith('.odt'):
        rows = parse_docx_timed_script(file_bytes)
    elif fn.endswith('.pdf'):
        rows = parse_pdf_timed_script(file_bytes)
    elif fn.endswith('.json'):
        rows = parse_json_timed_script(file_bytes)
    elif fn.endswith('.txt'):
        rows = parse_txt_timed_script(file_bytes)
    else:
        # Try docx first, then pdf, then txt
        try:
            rows = parse_docx_timed_script(file_bytes)
        except Exception:
            try:
                rows = parse_pdf_timed_script(file_bytes)
            except Exception:
                rows = parse_txt_timed_script(file_bytes)

    total_duration = "00:00"
    if rows:
        last_time = rows[-1]["time"]
        # Extract the end or single timestamp
        if "-" in last_time:
            total_duration = last_time.split("-")[-1].strip()
        else:
            total_duration = last_time

    title = filename.rsplit('.', 1)[0].replace('_', ' ').replace('-', ' ').title()

    return {
        "title": title,
        "total_rows": len(rows),
        "total_duration": total_duration,
        "rows": rows
    }
